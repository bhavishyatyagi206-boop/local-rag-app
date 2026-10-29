"""
Core RAG Engine: Extraction, Chunking, BM25 / Vector Hybrid Retrieval, and Extractive QA.
"""

from __future__ import annotations
import math
import os
import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union


class IngestError(Exception):
    """Raised when document extraction, parsing, or validation fails."""
    pass


@dataclass
class Chunk:
    id: str
    doc_id: str
    doc_name: str
    page_number: int
    chunk_index: int
    text: str
    token_count: int
    embedding: Optional[List[float]] = None
    embedding_model: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        return {
            "id": self.id,
            "doc_id": self.doc_id,
            "doc_name": self.doc_name,
            "page_number": self.page_number,
            "chunk_index": self.chunk_index,
            "text": self.text,
            "token_count": self.token_count,
            "embedding": self.embedding,
            "embedding_model": self.embedding_model,
        }

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> Chunk:
        return cls(
            id=data["id"],
            doc_id=data["doc_id"],
            doc_name=data["doc_name"],
            page_number=int(data["page_number"]),
            chunk_index=int(data["chunk_index"]),
            text=data["text"],
            token_count=int(data.get("token_count", len(data["text"].split()))),
            embedding=data.get("embedding"),
            embedding_model=data.get("embedding_model"),
        )


STOP_WORDS = {
    "a", "about", "above", "after", "again", "against", "all", "am", "an", "and",
    "any", "are", "as", "at", "be", "because", "been", "before", "being",
    "below", "between", "both", "but", "by", "can", "cannot", "could", "did", "do",
    "does", "doing", "down", "during", "each", "few", "for", "from", "further",
    "had", "has", "have", "having", "he", "her", "here", "hers", "herself", "him",
    "himself", "his", "how", "i", "if", "in", "into", "is", "it", "its",
    "itself", "me", "more", "most", "my", "myself", "no", "nor", "not", "of", "off",
    "on", "once", "only", "or", "other", "ought", "our", "ours", "ourselves", "out",
    "over", "own", "same", "she", "should", "so", "some", "such", "than", "that",
    "the", "their", "theirs", "them", "themselves", "then", "there", "these", "they",
    "this", "those", "through", "to", "too", "under", "until", "up", "very", "was",
    "we", "were", "what", "when", "where", "which", "while",
    "who", "whom", "why", "with", "would", "you", "your", "yours", "yourself",
}


def _tokenize(text: str) -> List[str]:
    """Lowercase alphanumeric tokenization."""
    return re.findall(r"\b[a-zA-Z0-9_]+\b", text.lower())


def extract(file_path: Union[str, Path], filename: Optional[str] = None) -> List[Tuple[int, str]]:
    """
    Extract text content from a file, returning a list of (page_number, text) tuples.
    Page numbers are 1-indexed.
    Supported extensions: .txt, .docx, .xlsx, .pdf
    Raises IngestError on parsing failure, corrupt file, unsupported format, or missing text.
    """
    path = Path(file_path)
    fname = filename or path.name
    ext = Path(fname).suffix.lower()

    if not path.is_file():
        raise IngestError(f"File not found or unreadable: {file_path}")

    if ext == ".txt":
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        except Exception as e:
            raise IngestError(f"Failed to read text file: {e}") from e

        if not content.strip():
            raise IngestError("Text file is empty")
        return [(1, content)]

    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            content = "\n\n".join(paragraphs).strip()
            if not content:
                raise IngestError("DOCX file contains no extractable text")
            return [(1, content)]
        except IngestError:
            raise
        except Exception as e:
            raise IngestError(f"Failed to read DOCX file: {e}") from e

    elif ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            lines = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                lines.append(f"--- Sheet: {sheetname} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
            content = "\n".join(lines).strip()
            if not content or content == "--- Sheet: " + (wb.sheetnames[0] if wb.sheetnames else "") + " ---":
                raise IngestError("XLSX file contains no extractable text")
            return [(1, content)]
        except IngestError:
            raise
        except Exception as e:
            raise IngestError(f"Failed to read XLSX file: {e}") from e

    elif ext == ".pdf":
        try:
            import fitz  # PyMuPDF
        except ImportError as e:
            raise IngestError("PyMuPDF (fitz) is required for PDF ingestion") from e

        try:
            doc = fitz.open(path)
        except Exception as e:
            raise IngestError(f"Corrupt PDF file or unable to open: {e}") from e

        try:
            if doc.is_encrypted:
                raise IngestError("Password-protected PDF")

            pages: List[Tuple[int, str]] = []
            total_text_len = 0
            for page_idx, page in enumerate(doc, start=1):
                page_text = page.get_text() or ""
                trimmed = page_text.strip()
                if trimmed:
                    pages.append((page_idx, trimmed))
                    total_text_len += len(trimmed)

            if total_text_len == 0 or not pages:
                raise IngestError("PDF has no text layer")

            return pages
        finally:
            doc.close()

    else:
        raise IngestError(f"Unsupported file format: '{ext}'. Allowed: .txt, .docx, .xlsx, .pdf")


def chunk_text(
    text: str,
    doc_id: str,
    doc_name: str,
    page_number: int = 1,
    start_chunk_idx: int = 0,
    chunk_size: int = 500,
    overlap: int = 50,
) -> List[Chunk]:
    """
    Split text into chunks with sliding window character count, preserving page numbers.
    """
    cleaned = text.strip()
    if not cleaned:
        return []

    if len(cleaned) <= chunk_size:
        return [
            Chunk(
                id=f"{doc_id}_p{page_number}_c{start_chunk_idx}",
                doc_id=doc_id,
                doc_name=doc_name,
                page_number=page_number,
                chunk_index=start_chunk_idx,
                text=cleaned,
                token_count=len(cleaned.split()),
            )
        ]

    chunks: List[Chunk] = []
    step = max(1, chunk_size - overlap)
    idx = start_chunk_idx

    start = 0
    while start < len(cleaned):
        end = min(start + chunk_size, len(cleaned))
        # If not at the end of the text, try to break at paragraph or sentence boundary
        if end < len(cleaned):
            break_pt = cleaned.rfind("\n\n", start, end)
            if break_pt == -1 or break_pt < start + (chunk_size // 2):
                break_pt = cleaned.rfind(". ", start, end)
            if break_pt != -1 and break_pt > start + (chunk_size // 3):
                end = break_pt + 1

        chunk_str = cleaned[start:end].strip()
        if chunk_str:
            chunks.append(
                Chunk(
                    id=f"{doc_id}_p{page_number}_c{idx}",
                    doc_id=doc_id,
                    doc_name=doc_name,
                    page_number=page_number,
                    chunk_index=idx,
                    text=chunk_str,
                    token_count=len(chunk_str.split()),
                )
            )
            idx += 1

        if end >= len(cleaned):
            break
        start += step

    return chunks


class BM25Index:
    """Okapi BM25 implementation for lexical keyword search."""

    def __init__(self, k1: float = 1.5, b: float = 0.75):
        self.k1 = k1
        self.b = b
        self.corpus: List[Chunk] = []
        self.doc_len: List[int] = []
        self.avgdl: float = 0.0
        self.doc_freqs: List[Dict[str, int]] = []
        self.idf: Dict[str, float] = {}

    def index(self, chunks: List[Chunk]) -> None:
        self.corpus = chunks
        self.doc_len = []
        self.doc_freqs = []
        df: Dict[str, int] = {}

        for chunk in chunks:
            tokens = _tokenize(chunk.text)
            self.doc_len.append(len(tokens))
            freq: Dict[str, int] = {}
            for t in tokens:
                freq[t] = freq.get(t, 0) + 1
            self.doc_freqs.append(freq)
            for t in freq:
                df[t] = df.get(t, 0) + 1

        n_docs = len(chunks)
        self.avgdl = sum(self.doc_len) / n_docs if n_docs > 0 else 0.0
        self.idf = {}
        for word, freq in df.items():
            # Standard Lucene/BM25 IDF with smoothing
            self.idf[word] = math.log(1.0 + (n_docs - freq + 0.5) / (freq + 0.5))

    def score(self, query: str) -> List[float]:
        q_tokens = _tokenize(query)
        scores = [0.0] * len(self.corpus)
        if not self.corpus or not q_tokens or self.avgdl == 0:
            return scores

        # Identify substantive content terms vs stop words
        content_tokens = [t for t in q_tokens if t not in STOP_WORDS]
        tokens_to_score = content_tokens if content_tokens else q_tokens

        for idx, freq_dict in enumerate(self.doc_freqs):
            dl = self.doc_len[idx]
            # If we have content terms, ensure at least one content term is present
            if content_tokens and not any(t in freq_dict for t in content_tokens):
                scores[idx] = 0.0
                continue

            doc_score = 0.0
            denom_part = self.k1 * (1.0 - self.b + self.b * (dl / self.avgdl))
            for q in tokens_to_score:
                if q in freq_dict:
                    f = freq_dict[q]
                    idf_val = self.idf.get(q, 0.0)
                    doc_score += idf_val * ((f * (self.k1 + 1.0)) / (f + denom_part))
            scores[idx] = max(0.0, doc_score)

        return scores


class EmbeddingProvider:
    """Base interface for dense vector embeddings."""

    name: str = "base"

    def embed_texts(self, texts: List[str]) -> List[List[float]]:
        raise NotImplementedError

    def embed_query(self, query: str) -> List[float]:
        return self.embed_texts([query])[0]


class OllamaEmbeddingProvider(EmbeddingProvider):
    """Dense vector embeddings via Ollama (e.g. nomic-embed-text)."""

    def __init__(self, model: str = "nomic-embed-text", base_url: str = "http://localhost:11434"):
        self.model = model
        self.base_url = base_url.rstrip("/")
        self.name = f"ollama:{model}"

    def embed_texts(self, texts: List[str]) -> List[List[float]]:
        import httpx
        embeddings: List[List[float]] = []
        with httpx.Client(timeout=30.0) as client:
            for text in texts:
                resp = client.post(
                    f"{self.base_url}/api/embeddings",
                    json={"model": self.model, "prompt": text},
                )
                if resp.status_code != 200:
                    raise RuntimeError(f"Ollama embedding failed ({resp.status_code}): {resp.text}")
                data = resp.json()
                embeddings.append(data["embedding"])
        return embeddings


class DeterministicLocalEmbeddingProvider(EmbeddingProvider):
    """
    Lightweight deterministic vector projection (64-dim) for testing and local hybrid RRF.
    Ensures vector storage, SQLite serialization, and RRF fusion work standalone.
    """

    def __init__(self, name: str = "local-hashing-v1", dim: int = 64):
        self.name = name
        self.dim = dim

    def embed_texts(self, texts: List[str]) -> List[List[float]]:
        import hashlib
        results: List[List[float]] = []
        for text in texts:
            vec = [0.0] * self.dim
            tokens = _tokenize(text)
            for t in tokens:
                h = int(hashlib.md5(t.encode("utf-8")).hexdigest(), 16)
                idx = h % self.dim
                sign = 1.0 if ((h >> 8) & 1) else -1.0
                vec[idx] += sign

            # L2 normalize
            norm = math.sqrt(sum(v * v for v in vec))
            if norm > 0:
                vec = [v / norm for v in vec]
            results.append(vec)
        return results


def cosine_similarity(vec_a: List[float], vec_b: List[float]) -> float:
    """Compute cosine similarity between two unit vectors."""
    if not vec_a or not vec_b or len(vec_a) != len(vec_b):
        return 0.0
    dot = sum(a * b for a, b in zip(vec_a, vec_b))
    return max(-1.0, min(1.0, dot))


class KnowledgeBase:
    """In-memory Knowledge Base with BM25 and optional dense vector retrieval."""

    def __init__(self, id: Optional[str] = None, name: str = "Default KB"):
        self.id = id or str(uuid.uuid4())
        self.name = name
        self.documents: Dict[str, Dict[str, Any]] = {}
        self.chunks: List[Chunk] = []
        self.bm25 = BM25Index()
        self.embedding_provider: Optional[EmbeddingProvider] = None

    def add_document(
        self,
        doc_id: str,
        name: str,
        pages: List[Tuple[int, str]],
        metadata: Optional[Dict[str, Any]] = None,
        chunk_size: int = 500,
        overlap: int = 50,
    ) -> List[Chunk]:
        """Index a document given pre-extracted pages."""
        # Remove any prior chunks for this doc_id
        self.remove_document(doc_id)

        doc_chunks: List[Chunk] = []
        global_chunk_idx = 0
        for page_num, text in pages:
            page_chunks = chunk_text(
                text=text,
                doc_id=doc_id,
                doc_name=name,
                page_number=page_num,
                start_chunk_idx=global_chunk_idx,
                chunk_size=chunk_size,
                overlap=overlap,
            )
            doc_chunks.extend(page_chunks)
            global_chunk_idx += len(page_chunks)

        self.documents[doc_id] = {
            "id": doc_id,
            "name": name,
            "page_count": len(pages),
            "chunk_count": len(doc_chunks),
            "metadata": metadata or {},
        }
        self.chunks.extend(doc_chunks)
        self.reindex()
        return doc_chunks

    def remove_document(self, doc_id: str) -> None:
        if doc_id in self.documents:
            del self.documents[doc_id]
        self.chunks = [c for c in self.chunks if c.doc_id != doc_id]
        self.reindex()

    def reindex(self) -> None:
        self.bm25.index(self.chunks)

    def set_embedding_provider(self, provider: Optional[EmbeddingProvider]) -> None:
        """Assign embedding provider and trigger re-indexing if model changes."""
        self.embedding_provider = provider
        if provider:
            self.ensure_embeddings()

    def ensure_embeddings(self) -> None:
        """Re-embed chunks if embedding model changed or embeddings are missing."""
        if not self.embedding_provider or not self.chunks:
            return

        target_model = self.embedding_provider.name
        needs_reembed = any(c.embedding is None or c.embedding_model != target_model for c in self.chunks)
        if needs_reembed:
            texts = [c.text for c in self.chunks]
            vectors = self.embedding_provider.embed_texts(texts)
            for chunk, vec in zip(self.chunks, vectors):
                chunk.embedding = vec
                chunk.embedding_model = target_model

    def retrieve(
        self,
        query: str,
        top_k: int = 5,
        doc_ids: Optional[List[str]] = None,
    ) -> List[Tuple[Chunk, float]]:
        """
        Retrieve chunks matching query using BM25 or Hybrid RRF (Reciprocal Rank Fusion).
        Enforces cross-document filtering and model consistency.
        """
        if not self.chunks:
            return []

        allowed_docs = set(doc_ids) if doc_ids is not None else None
        candidate_chunks = [c for c in self.chunks if allowed_docs is None or c.doc_id in allowed_docs]
        if not candidate_chunks:
            return []

        # 1. BM25 scoring
        bm25_scores = self.bm25.score(query)
        bm25_scored: List[Tuple[Chunk, float]] = []
        for chunk, score in zip(self.chunks, bm25_scores):
            if allowed_docs is None or chunk.doc_id in allowed_docs:
                bm25_scored.append((chunk, float(score)))

        # If no dense embedding provider, return BM25 results
        if not self.embedding_provider:
            bm25_scored.sort(key=lambda x: x[1], reverse=True)
            return bm25_scored[:top_k]

        # 2. Hybrid Dense + Sparse via Reciprocal Rank Fusion (RRF)
        self.ensure_embeddings()
        try:
            q_vec = self.embedding_provider.embed_query(query)
        except Exception:
            # Fallback to BM25 if embedding provider fails
            bm25_scored.sort(key=lambda x: x[1], reverse=True)
            return bm25_scored[:top_k]

        # Rank BM25
        bm25_scored.sort(key=lambda x: x[1], reverse=True)
        bm25_rank_map = {chunk.id: rank for rank, (chunk, score) in enumerate(bm25_scored, start=1)}

        # Rank Dense
        dense_scored: List[Tuple[Chunk, float]] = []
        for chunk in candidate_chunks:
            if chunk.embedding and chunk.embedding_model == self.embedding_provider.name:
                sim = cosine_similarity(q_vec, chunk.embedding)
            else:
                sim = 0.0
            dense_scored.append((chunk, sim))

        dense_scored.sort(key=lambda x: x[1], reverse=True)
        dense_rank_map = {chunk.id: rank for rank, (chunk, _) in enumerate(dense_scored, start=1)}

        # Fuse with Reciprocal Rank Fusion (RRF) formula: 1 / (60 + rank)
        fused: List[Tuple[Chunk, float]] = []
        k_const = 60
        for chunk in candidate_chunks:
            bm25_rank = bm25_rank_map.get(chunk.id, len(candidate_chunks) + 1)
            dense_rank = dense_rank_map.get(chunk.id, len(candidate_chunks) + 1)
            
            bm25_score = next((s for c, s in bm25_scored if c.id == chunk.id), 0.0)
            dense_sim = next((s for c, s in dense_scored if c.id == chunk.id), 0.0)

            # If neither signal shows relevance, skip
            if bm25_score <= 0.0 and dense_sim < 0.35:
                continue

            rrf_score = (1.0 / (k_const + bm25_rank)) + (1.0 / (k_const + dense_rank))
            fused.append((chunk, float(rrf_score)))

        fused.sort(key=lambda x: x[1], reverse=True)
        return fused[:top_k]

    def ask(
        self,
        query: str,
        top_k: int = 5,
        doc_ids: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """
        Extractive QA fallback: selects the best matching evidence from retrieved chunks.
        Never fabricates answers or citations.
        """
        retrieved = self.retrieve(query=query, top_k=top_k, doc_ids=doc_ids)
        if not retrieved or all(score <= 0.0 for _, score in retrieved):
            return {
                "answer": "No relevant context found in the selected documents for this question.",
                "status": "no_relevant_context",
                "citations": [],
                "debug": {
                    "query": query,
                    "retrieved_chunks": [],
                    "scores": [],
                },
            }

        # Filter out 0-score chunks
        relevant = [(c, s) for c, s in retrieved if s > 0.0]
        if not relevant:
            relevant = retrieved[:1]

        best_chunk, top_score = relevant[0]

        # Extract the most relevant sentence from best_chunk
        q_tokens = set(_tokenize(query))
        sentences = re.split(r"(?<=[.!?])\s+", best_chunk.text)
        best_sentence = best_chunk.text
        max_overlap = -1

        for sent in sentences:
            sent_tokens = set(_tokenize(sent))
            overlap_cnt = len(q_tokens.intersection(sent_tokens))
            if overlap_cnt > max_overlap:
                max_overlap = overlap_cnt
                best_sentence = sent.strip()

        # Build citations
        citations = []
        seen_chunks = set()
        for chunk, score in relevant:
            if chunk.id in seen_chunks:
                continue
            seen_chunks.add(chunk.id)
            excerpt = chunk.text[:250].strip() + ("..." if len(chunk.text) > 250 else "")
            citations.append({
                "chunk_id": chunk.id,
                "doc_id": chunk.doc_id,
                "doc_name": chunk.doc_name,
                "page_number": chunk.page_number,
                "score": round(score, 4),
                "excerpt": excerpt,
            })

        answer = (
            f"Based on {best_chunk.doc_name} (Page {best_chunk.page_number}):\n\n"
            f"{best_sentence}"
        )

        return {
            "answer": answer,
            "status": "answered",
            "citations": citations,
            "debug": {
                "query": query,
                "retrieved_chunks": [c.id for c, _ in relevant],
                "scores": [round(s, 4) for _, s in relevant],
            },
        }
